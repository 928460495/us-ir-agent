"""相对估值接线与工作簿页的美股版（计划 Task 8）。用 us-dcf 的 AAPL 工作簿固件。不联网。"""

from __future__ import annotations

import shutil
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from tests.test_relval import _peers_pe, make_result

AAPL_WB = Path(__file__).resolve().parent.parent / "us-dcf" / "tests" / "fixtures" / "AAPL_DCF.xlsx"
DCF_OK = SimpleNamespace(scenario_ps={"bear": D("95"), "base": D("134"), "bull": D("174")},
                         mc={"p5": D("80"), "p95": D("190")})


def _us_res():
    r = make_result()
    r.target.code = "AAPL"
    r.target.eps_ntm, r.target.revenue = D("9.1"), D("416161000000")
    r.comps["pe_ntm"] = r.comps["pe"]
    r.comps_peers = [p for p in _peers_pe(10, 20, 30, 40, 50, 60)]
    for p in r.comps_peers:
        p.price, p.eps_ntm, p.ntm_basis = D(30), D(2), "NTM：FY2027 × 0.921 + FY2028 × 0.079"
    return r


def _text(path):
    ws = load_workbook(path)["9_相对估值"]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


@pytest.fixture
def us_wb(tmp_path):
    p = tmp_path / "AAPL.xlsx"
    shutil.copy(AAPL_WB, p)
    return str(p)


def test_美股相对估值页单位是百万美元与美元(us_wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(us_wb, _us_res(), DCF_OK)
    t = _text(us_wb)
    assert "百万美元" in t and "百万股" in t and "（美元）" in t
    assert "亿元" not in t and "（元）" not in t and "A 股折旧摊销" not in t


def test_美股可比表有PE_NTM一列与NTM口径_隐含价公式覆盖新倍数(us_wb):
    from ir_agent.relval_sheet import write_sheet
    write_sheet(us_wb, _us_res(), DCF_OK)
    ws = load_workbook(us_wb)["9_相对估值"]
    t = _text(us_wb)
    assert "PE(NTM)" in t and "NTM 口径" in t and "NTM EPS" in t
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) >= 12                                  # 四个倍数 × P25 / 中位数 / P75


def test_美股走gather_us并按市场选倍数(monkeypatch):
    from ir_agent import full, relval, relval_data, relval_data_us
    seen = {}
    tgt = make_result().target
    tgt.code = "AAPL"

    def fake_us(r, as_of, snap):
        seen["us"] = True
        return tgt, [], [], D(1), D(1), []
    monkeypatch.setattr(relval_data_us, "gather_us", fake_us)
    monkeypatch.setattr(relval_data, "gather", lambda *a, **k: pytest.fail("美股不该走 A 股 gather"))
    real = relval.comps_bands

    def spy(peers, target, shares, kinds=relval.KINDS):
        seen["kinds"] = kinds
        return real(peers, target, shares, kinds=kinds)
    monkeypatch.setattr(relval, "comps_bands", spy)
    r = SimpleNamespace(code="AAPL", peer_set=SimpleNamespace(classification=SimpleNamespace(used_path="FF49-35")))
    res, _sk = full._relative_valuation("AAPL", r, date(2026, 10, 26), None, None)
    assert seen["us"] and "pe_ntm" in seen["kinds"] and res is not None


def test_美股估值结论写美元与us_dcf():
    from ir_agent import full
    assert full._value_reason(SimpleNamespace(value_per_share=D("134.78"), price=D("342.8"), upside=D("-0.6"),
                                              wacc=D("0.09"), terminal_share=D("0.6")),
                              "谢海量", "AAPL").count("美元") == 2
    cn = full._value_reason(SimpleNamespace(value_per_share=D("30"), price=D("25"), upside=D("0.2"),
                                            wacc=D("0.09"), terminal_share=D("0.6")), "x", "600183")
    assert "元" in cn and "美元" not in cn
    assert full._engine_name("AAPL") == "us-dcf" and full._engine_name("600183") == "ashare-dcf"
