"""美股对账：us-dcf 工作簿「1b_驱动分解」的合成格 ≡ ir-agent 计算引擎（与 A 股 test_formula_workbook 同一套断言）。

工作簿合成格用 us_verify 自己的单元格求值（与公式引擎验算同一套），两边各算各的。
"""
from __future__ import annotations

import importlib.util
import sys
from decimal import Decimal as D
from pathlib import Path

import pytest

from ir_agent import assume
from ir_agent.assume import Case
from ir_agent.methodology import parse as parse_meth

US_DCF = Path(__file__).resolve().parent.parent / "us-dcf"
sys.path.insert(0, str(US_DCF))
import us_build as B  # noqa: E402
import us_verify as V  # noqa: E402

_spec = importlib.util.spec_from_file_location("us_dcf_conftest", US_DCF / "tests" / "conftest.py")
_cf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_cf)
Q = "「依据原文」"


def _rows(spec: dict, years: int) -> str:
    def v(x, y):
        return x(y) if callable(x) else x
    return "\n".join(f"假设: {k} | 第{y}年 | {v(lo, y)} | {v(mid, y)} | {v(hi, y)} | 理由 {Q}"
                     for y in range(1, years + 1) for k, (lo, mid, hi) in spec.items())


LEVEL = parse_meth("""模式: 水平
因子: 量 | iPhone 出货量 | 百万台 | 10-K 业务 | 出货量
因子: 价 | iPhone 平均售价 | 美元/台 | 业绩电话会 | 平均售价
因子: 本 | iPhone 单位成本 | 美元/台 | 10-K 成本 | 单位成本
公式: 营收 = 量 × 价
公式: 毛利率 = (价 − 本) ÷ 价
""").model
LEVEL_ROWS = {"量": tuple((lambda g: lambda y: f"{230 * (1 + g) ** y:.4f}")(g) for g in (0.0, 0.03, 0.06)),
              "价": ("900", "950", "1000"), "本": ("540", "530", "520")}
LEVEL_BASE = f"基期: 量 | 230 | {Q}\n基期: 价 | 920 | {Q}\n基期: 本 | 535 | {Q}"
CHANGE = parse_meth("""模式: 变化
因子: 量增 | 出货量同比变化率 | 变化率 | 10-K | 出货量
因子: 价增 | 平均售价同比变化率 | 变化率 | 电话会 | 平均售价
因子: 结构 | 服务占比提升对毛利率的贡献 | pp | 10-K | 服务收入
公式: 营收 = 上期(营收) × (1 + 量增) × (1 + 价增)
公式: 毛利率 = 上期(毛利率) + 结构 / 100
""").model
CHANGE_ROWS = {"量增": ("-3%", "2%", "5%"), "价增": ("-2%", "1%", "3%"), "结构": ("0", "0.4", "0.8")}


def _build(model, rows, extra, tmp_path):
    cal = _cf.load_cal()
    years = len(cal["periods"]["proj"])
    a = assume.parse(_rows(rows, years) + ("\n" + extra if extra else ""), years, decomps=model)
    a.labels = list(cal["periods"]["proj"])
    out = assume.apply_to(cal, a)
    out["judg"]["drivers"] = assume.to_driver_block(a, cal)
    path = tmp_path / "w.xlsx"
    B.build(out, str(path))
    return a, cal, path


def _sheet_values(path, suffix, cal):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ds = wb["1b_驱动分解"]
    rowmap = {str(ds.cell(r, 1).value): r for r in range(1, ds.max_row + 1)}
    out = {case: [D(str(V._cell_value(wb, ds, f"{col}{rowmap[f'{y} {suffix}']}")))
                  for y in cal["periods"]["proj"]]
           for col, case in zip("BCD", ("bull", "base", "bear"))}
    return out, wb, ds, rowmap


def _engine(a, cal):
    bm = assume._base_margin_of(cal)
    bases = assume.margin_bases(a, bm, assume.anchor_from_cal(cal))
    return {c.name.lower(): assume.synthesize(a, c, bm, bases) for c in (Case.BULL, Case.BASE, Case.BEAR)}


@pytest.mark.parametrize("model,rows,extra", [(LEVEL, LEVEL_ROWS, LEVEL_BASE), (CHANGE, CHANGE_ROWS, "")],
                         ids=["水平模式", "变化模式_上期毛利率"])
def test_工作簿合成格与ir_agent引擎逐档逐年一致(model, rows, extra, tmp_path):
    a, cal, path = _build(model, rows, extra, tmp_path)
    want = _engine(a, cal)
    growth, *_ = _sheet_values(path, "营收增速（合成）", cal)
    margin, *_ = _sheet_values(path, "毛利率（合成）", cal)
    for case in ("bull", "base", "bear"):
        for y in range(len(growth[case])):
            assert abs(growth[case][y] - want[case][0][y]) < D("1e-9"), (case, y)
            assert abs(margin[case][y] - want[case][1][y]) < D("1e-9"), (case, y)


def test_工作簿公式引擎验算全过_三档价值分化(tmp_path):
    _, _, path = _build(LEVEL, LEVEL_ROWS, LEVEL_BASE, tmp_path)
    sc = {}
    assert V.verify_all_scenarios(str(path), quiet=True, values=sc) == []
    assert sc["Bull"] > sc["Base"] > sc["Bear"]


def test_假设页增速是跨表引用_改一个因子合成格随之变化(tmp_path):
    _, cal, path = _build(CHANGE, CHANGE_ROWS, "", tmp_path)
    before, wb, ds, rowmap = _sheet_values(path, "营收增速（合成）", cal)
    assert "基期 营收（百万美元）" in rowmap
    a = wb["1_Assumptions"]
    y2 = cal["periods"]["proj"][1]
    g_row = next(r for r in range(1, a.max_row + 1) if a.cell(r, 1).value == f"{y2} 营收增速")
    assert str(a.cell(g_row, 3).value).startswith("='1b_驱动分解'!")
    ds[f"C{rowmap[f'{y2} 平均售价同比变化率']}"] = 0.20
    after = D(str(V._cell_value(wb, ds, f"C{rowmap[f'{y2} 营收增速（合成）']}")))
    assert after != before["base"][1]


def test_read_inputs读到驱动表合成出的增速(tmp_path):
    a, cal, path = _build(CHANGE, CHANGE_ROWS, "", tmp_path)
    want = _engine(a, cal)["base"][0]
    got = V.read_inputs(str(path))["growth"]
    # 第 1 年被季度锚定覆盖（三档相同），从第 2 年起应等于合成增速
    assert all(abs(D(str(g)) - w) < D("1e-9") for g, w in zip(got[1:], want[1:]))
