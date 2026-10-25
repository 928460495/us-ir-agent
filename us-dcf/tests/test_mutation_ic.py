# -*- coding: utf-8 -*-
"""变异测试：破坏新的 IC / 研发资本 / 终值 ROIC 链路，确认验证器会报错。"""
import sys, shutil, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
import us_verify as V
SRC = os.environ.get("US_TEST_XLSX") or sys.argv[1]
def rm(ws):
    m={}
    for r in range(1,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and v.strip(): m.setdefault(v.strip(),r)
    return m
def mut(tag, fn):
    d=f"/tmp/_mic_{tag}.xlsx"; shutil.copy(SRC,d)
    wb=load_workbook(d); fn(wb); wb.save(d)
    fails,_=V.verify(d,quiet=True)
    ok=(len(fails)==0) if tag.startswith("00") else (len(fails)>0)
    print(f"{tag:36s} → {len(fails):3d} 项失败  {'✅' if ok else '❌ 漏网'}")
    if fails and not tag.startswith("00"): print("      例:",fails[0])

mut("00_不破坏（对照组）", lambda wb: None)
# IC 少算商誉无形 + 其他非流动
mut("01_IC 退回 NWC+PP&E", lambda wb:(lambda w,m: w.cell(m["投入资本 IC（申报口径）"],2,
    f"=B{m['FY2030E 营运资本 NWC']}+B{m['FY2030E 固定资产净额']}"))(wb["5_DCF"],rm(wb["5_DCF"])))
# 研发资本开关被忽略
mut("02_生效IC 忽略研发资本", lambda wb:(lambda w,m: w.cell(m["生效投入资本 IC"],2,
    f"=B{m['投入资本 IC（申报口径）']}"))(wb["5_DCF"],rm(wb["5_DCF"])))
# 研发资本滚动漏掉本期研发
mut("03_研发资本漏本期支出", lambda wb: wb["3_Balance"].__setitem__("H33","='3_Balance'!G33-'3_Balance'!H32"))
# 研发摊销用错年限
mut("04_研发摊销年限改错", lambda wb: wb["3_Balance"].__setitem__("H32","='3_Balance'!G33/99"))
# 终值 ROIC 生效值忽略模式开关
mut("05_终值ROIC 忽略模式开关", lambda wb:(lambda w,m: w.cell(m["终值 ROIC（生效）"],2,
    f"=B{m['终值 ROIC（收敛口径）']}"))(wb["5_DCF"],rm(wb["5_DCF"])))
