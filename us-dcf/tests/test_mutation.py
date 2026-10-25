# -*- coding: utf-8 -*-
"""变异测试：刻意破坏每条关键接线，确认验证器真的会报错。"""
import os, sys, shutil, warnings; warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import load_workbook
import us_verify as V
SRC = os.environ.get("US_TEST_XLSX") or sys.argv[1]
os.makedirs("/tmp/us_mut", exist_ok=True)

def rowmap(ws):
    m = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip(): m.setdefault(a.strip(), r)
    return m

def mut(tag, fn):
    dst = f"/tmp/us_mut/{tag}.xlsx"; shutil.copy(SRC, dst)
    wb = load_workbook(dst); fn(wb); wb.save(dst)
    fails, _ = V.verify(dst, quiet=True)
    ok = (len(fails) == 0) if tag.startswith("00") else (len(fails) > 0)
    print(f"{tag:32s} → {len(fails):3d} 项失败  {'✅' if ok else '❌'}")
    if fails and not tag.startswith("00"): print("      例:", fails[0])

mut("00_不破坏（对照组）", lambda wb: None)
# 注：不要用「漏掉其他经营损益」做变异 —— 该残差项对某些公司恰为 0（AAPL 就是），
# 变异会退化成空操作，看起来像验证器失灵，其实是变异本身无效。改用必然非零的 SG&A。
mut("01_EBIT漏SG&A", lambda wb: wb["2_Income"].__setitem__("G14", "=G9-G11+G13"))
mut("02_折旧率用错期末PPE", lambda wb: wb["4_CashFlow"].__setitem__(
        "G7", "='3_Balance'!G12*'1_Assumptions'!$B$44"))
mut("03_FCFF误加回SBC", lambda wb: wb["4_CashFlow"].__setitem__(
        "G19", "='2_Income'!G14*(1-'1_Assumptions'!$B$36)+G7+G8+G9+G16+SUM(G10:G14)"))
mut("04_资产总计漏长期投资", lambda wb: wb["3_Balance"].__setitem__("G16", "=G11+G12+G14+G15"))
mut("05_权益漏SBC", lambda wb: wb["3_Balance"].__setitem__(
        "G25", "='3_Balance'!F25+'2_Income'!G22+'4_CashFlow'!G20+'4_CashFlow'!G21+'4_CashFlow'!G22"))
mut("06_股数漏回购", lambda wb: wb["4_CashFlow"].__setitem__(
        "G30", "='1_Assumptions'!$B$71+G29"))
mut("07_终值漏降级保护", lambda wb: (lambda w, m: w.cell(
        m["终值 TV"], 2, f"=B{m['终值 FCFF']}/(B{m['WACC']}-B{m['永续增长率 g']})"))(
        wb["5_DCF"], rowmap(wb["5_DCF"])))
mut("08_净负债误加租赁负债", lambda wb: (lambda w, m: w.cell(
        m["减：净负债（基期）"], 2,
        "='1_Assumptions'!B71+'1_Assumptions'!B72"))(wb["5_DCF"], rowmap(wb["5_DCF"])))
