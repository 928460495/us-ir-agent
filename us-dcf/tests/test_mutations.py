"""变异测试：刻意破坏一条关键接线，验证器必须报错；对照组必须 0 失败。

由 test_mutation.py / test_mutation_ic.py（导入即执行的脚本）改写为 pytest 用例：
- 假设页一律按**标签**定位。旧脚本写死的 '1_Assumptions'!$B$36 / $B$44 / $B$71 在旧工作簿里
  就已错位（分别指向「全期毛利率 Δ」「占比法隐含营收」「存货」），变异只是碰巧改坏了公式，
  证明不了验证器抓得住那个特定错误；
- 报表页（2_Income / 3_Balance / 4_CashFlow）的行号新旧一致（已逐行核对），沿用；
- 「终值漏降级保护」对 AAPL 是空操作（ROIC 远高于 WACC，从不降级），改为先把终值 ROIC
  设成「收敛至 WACC − 1%」强制降级，再删保护；另设只改设置不删保护的对照组。
"""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import us_verify as V

FIX = Path(__file__).parent / "fixtures" / "AAPL_DCF.xlsx"


def _rows(ws) -> dict:
    m = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            m.setdefault(v.strip(), r)
    return m


def _a(wb, label: str) -> str:
    """假设页某个标签的 B 列绝对引用。"""
    return f"'1_Assumptions'!$B${_rows(wb['1_Assumptions'])[label]}"


def _set(sheet, label, value):
    def f(wb):
        ws = wb[sheet]
        ws.cell(_rows(ws)[label], 2).value = value
    return f


def _force_degrade(wb):
    _set("1_Assumptions", "终值 ROIC 处理（1=沿用末年 2=收敛至 WACC+溢价）", 2)(wb)
    _set("1_Assumptions", "终值 ROIC 收敛溢价（模式2 用）", -0.01)(wb)


def _append(sheet, addr, extra):
    def f(wb):
        wb[sheet][addr] = str(wb[sheet][addr].value) + extra
    return f


def _last_proj(wb) -> str:
    hdr = [wb["2_Income"].cell(5, c).value for c in range(2, 60)]
    return [h for h in hdr if h and str(h).endswith("E")][-1]


def _drop_degrade(wb):
    _force_degrade(wb)
    d = wb["5_DCF"]; m = _rows(d)
    d.cell(m["终值 TV"], 2).value = f"=B{m['终值 FCFF']}/(B{m['WACC']}-B{m['永续增长率 g']})"


def _ic_back_to_nwc_ppe(wb):
    d = wb["5_DCF"]; m = _rows(d); y = _last_proj(wb)
    d.cell(m["投入资本 IC（申报口径）"], 2).value = f"=B{m[f'{y} 营运资本 NWC']}+B{m[f'{y} 固定资产净额']}"


MUTATIONS = {
    "EBIT漏SG&A": lambda wb: wb["2_Income"].__setitem__("G14", "=G9-G11+G13"),
    "折旧率用错期末PPE": lambda wb: wb["4_CashFlow"].__setitem__(
        "G7", f"='3_Balance'!G12*{_a(wb, '折旧率（占期初固定资产净额）')}"),
    "FCFF误加回SBC": _append("4_CashFlow", "G19", "+G9"),
    "资产总计漏长期投资": lambda wb: wb["3_Balance"].__setitem__("G16", "=G11+G12+G14+G15"),
    "权益漏SBC": lambda wb: wb["3_Balance"].__setitem__(
        "G25", "='3_Balance'!F25+'2_Income'!G22+'4_CashFlow'!G20+'4_CashFlow'!G21+'4_CashFlow'!G22"),
    "股数漏回购": lambda wb: wb["4_CashFlow"].__setitem__("G30", f"={_a(wb, '摊薄股数（百万股）')}+G29"),
    "终值漏降级保护": _drop_degrade,
    "净负债误加租赁负债": lambda wb: _append(
        "5_DCF", f"B{_rows(wb['5_DCF'])['减：净负债（基期）']}",
        f"+{_a(wb, '（备查）经营租赁负债 — 非流动')}")(wb),
    "IC退回NWC加PPE": _ic_back_to_nwc_ppe,
    "生效IC忽略研发资本": lambda wb: (lambda d, m: d.cell(m["生效投入资本 IC"], 2, f"=B{m['投入资本 IC（申报口径）']}"))(
        wb["5_DCF"], _rows(wb["5_DCF"])),
    "研发资本漏本期支出": lambda wb: wb["3_Balance"].__setitem__("H33", "='3_Balance'!G33-'3_Balance'!H32"),
    "研发摊销年限改错": lambda wb: wb["3_Balance"].__setitem__("H32", "='3_Balance'!G33/99"),
    "终值ROIC忽略模式开关": lambda wb: (lambda d, m: d.cell(m["终值 ROIC（生效）"], 2, f"=B{m['终值 ROIC（收敛口径）']}"))(
        wb["5_DCF"], _rows(wb["5_DCF"])),
}


def _verify_after(tmp_path, fn):
    dst = tmp_path / "m.xlsx"
    shutil.copy(FIX, dst)
    wb = load_workbook(dst)
    fn(wb)
    wb.save(dst)
    return V.verify(str(dst), quiet=True)


@pytest.mark.parametrize("ctrl", [lambda wb: None, _force_degrade], ids=["原样", "强制降级"])
def test_对照组零失败(tmp_path, ctrl):
    fails, dv = _verify_after(tmp_path, ctrl)
    assert fails == []
    if ctrl is _force_degrade:
        assert dv["degrade"] == 1.0                     # 设置确实让终值进入降级区


@pytest.mark.parametrize("name", list(MUTATIONS))
def test_变异被验证器抓到(tmp_path, name):
    fails, _ = _verify_after(tmp_path, MUTATIONS[name])
    assert fails, f"变异「{name}」没有被验证器发现"
