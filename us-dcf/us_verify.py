# -*- coding: utf-8 -*-
"""美股 DCF 模板 — 验证层。

两件事：
  1. 用 formulas 引擎【真正跑一遍】工作簿里的公式；
  2. 在 Python 侧【完全独立地】把模型重算一遍，逐格比对。

只有两条独立路径得到同一个数，才算验过。只检查「模型内部是否自洽」是不够的 ——
自洽地错，是最难发现的错。
"""
from __future__ import annotations
import math, os, re as _re
import numpy as np
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL

TOL = 1e-6


def _rowmap(ws) -> dict:
    m = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            m.setdefault(v.strip(), r)
    return m


def _find(m: dict, key: str) -> int:
    """先精确匹配，再前缀匹配 —— 前缀优先会让「存货」命中「存货周转天数」。"""
    if key in m:
        return m[key]
    hits = [k for k in m if k.startswith(key)]
    if not hits:
        raise KeyError(key)
    return m[sorted(hits, key=len)[0]]


# 跨表引用解析 —— 自 ashare_verify 原样移植。「1b_驱动分解」接入后，1_Assumptions 的营收增速 /
# 毛利率三档格是 ='1b_驱动分解'!B12 这类跨表引用；不解析就会把公式字符串当值，独立重算整个断掉。
_REF_RE = re.compile(
    # 「!」不设为可选——本表引用已由 _SELF_REF_RE 先行拦截；这里若把「!」设为
    # 可选，形如 "=B1+C1" 的普通公式会被表名分组贪婪吞掉 "B1+" 再回溯匹配出
    # "C1"，被误判成引用。跨表引用必须显式带「!」，因此收紧为必选。
    r"^=\s*(?:'([^']+)'|([^!'\s]+))!\$?([A-Z]+)\$?(\d+)\s*$")
_SELF_REF_RE = re.compile(r"^=\s*\$?([A-Z]+)\$?(\d+)\s*$")
# 引用层数上限。原为 5；ir-agent 公式化（2026-10-14）后「1b_驱动分解」的营收
# 合成格逐年引用上一年（第 t 年 → 第 t−1 年 → … → 基期），加上 1_假设 → 增速行
# 的两跳，十年预测的链深约 2×10+3。64 仍能拦住真正的循环引用。
_MAX_REF_DEPTH = 64
_ARITH_CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d{1,7})")


class _CircularRefError(ValueError):
    """_cell_value 的引用层数守卫专用异常——是 ValueError 的子类（不改变既有
    `except ValueError` 调用方的行为），只为了在 _eval_arith 里能把它跟『解析
    失败』的普通 ValueError 区分开，见 _eval_arith 里的说明。"""


def _eval_arith(wb, sheet, expr, _depth):
    """对『仅含本表单元格引用与四则运算』的简单算式求值。

    驱动因子分解的『合成』行故意写成 Excel 公式（如 (1+C4)*(1+C5)-1 或
    0.1415+(C10+C11)/100）而非常量——这样改一个因子能顺着公式一路传导到
    2_利润表/5_DCF。但 read_inputs 的独立复算得先从这类公式里拿到一个
    浮点数：_resolve_ref 只认『整格就是一个引用』（=B48 / =1b_驱动分解!B5），
    对这种四则运算公式无能为力，会直接把公式字符串丢回去，下游 float()
    立刻炸掉。这里把公式里出现的单元格逐个换成其（可能还要跨表/递归解析
    出来的）数值，再对剩下的纯算式求值。只覆盖本项目自己生成的这类简单
    算式——不支持函数调用等更复杂形式；解析不出来就返回 None，调用方按
    老规矩原样返回公式字符串，不影响其余既有行为。
    """
    def repl(m):
        v = _cell_value(wb, sheet, f"{m.group(1)}{m.group(2)}", _depth + 1)
        return f"({float(v)})"
    try:
        text = _ARITH_CELL_RE.sub(repl, expr)
    except _CircularRefError:
        # 循环引用的诊断信息必须穿透这里，不能被下面那条兜底 except 吞掉——
        # 否则排查者看到的只是一个莫名其妙的 float() 类型错误，而不是「这里可能
        # 有循环引用」。别的 ValueError（比如 repl 里 float(v) 碰到解析不出数值
        # 的东西）仍然维持老行为：走 return None，让调用方原样吐回公式串。
        raise
    except (TypeError, ValueError):
        return None
    if not re.fullmatch(r"[0-9+\-*/().\s]+", text):
        return None          # 残留非数字字符（如函数名）——不是我们支持的算式
    if "**" in text:
        # 显式拒绝乘方：白名单按字符放行 '*'，两个相邻就拼成 Python 的乘方
        # 运算符。Excel 自己的乘方符号是 '^'（已被上面的 fullmatch 挡掉），
        # 只有手改过的工作簿才写得出 9**9**9 这种指数塔 —— 它会算出一个几亿位
        # 的大整数，把校验进程拖死。
        return None
    try:
        return float(eval(text, {"__builtins__": {}}, {}))       # noqa: S307
    except Exception:
        return None


def _resolve_ref(v):
    """把单元格里的引用式拆成 (表名 或 None, 地址)。不是引用则返回 None。

    此前只认本表引用（=B48）。接入驱动因子分解后，「1_假设」的增速与毛利率
    单元格要指向「1b_驱动分解」，不支持跨表就会让 read_inputs 拿到公式字符串
    然后 float() 抛异常 —— **独立复算验证会整个断掉**。
    """
    if not isinstance(v, str):
        return None
    m = _SELF_REF_RE.match(v.strip())
    if m:
        return None, f"{m.group(1)}{m.group(2)}"
    m = _REF_RE.match(v.strip())
    if not m:
        return None
    sheet = m.group(1) or m.group(2)
    if not sheet:
        return None
    return sheet, f"{m.group(3)}{m.group(4)}"


def _cell_value(wb, sheet, addr, _depth=0):
    """读一个单元格的值，逐层解析引用（含跨表）。"""
    if _depth > _MAX_REF_DEPTH:
        raise _CircularRefError(
            f"引用层数超过 {_MAX_REF_DEPTH} —— 可能存在循环引用")
    v = sheet[addr].value
    ref = _resolve_ref(v)
    if ref is None:
        if isinstance(v, str) and v.strip().startswith("="):
            ev = _eval_arith(wb, sheet, v.strip()[1:], _depth)
            if ev is not None:
                return ev
        return v
    name, target = ref
    nxt = wb[name] if name else sheet
    return _cell_value(wb, nxt, target, _depth + 1)


def read_inputs(path: str) -> dict:
    wb = load_workbook(path)
    a = wb["1_Assumptions"]; am = _rowmap(a)

    def _cell(ref, _d=0):
        """解析本表单元格，支持一层算式（如 =B116+B123）。
        ERP 那一格是活公式（成熟基准 + 地域加权 CRP），只认单格引用会读到字符串。"""
        v = a[ref].value
        return _expr(v, _d) if isinstance(v, str) else v

    def _expr(v, _d=0):
        if not isinstance(v, str) or not v.startswith("=") or _d > 5:
            return v
        body = v[1:].strip()
        if not _re.fullmatch(r"[\$A-Z0-9\.\+\-\*/\(\) ]+", body):
            return v                       # 含函数/跨表引用 → 不解析，交给上层
        def sub(mo):
            x = _cell(f"{mo.group(1)}{mo.group(2)}", _d + 1)
            return f"({float(x)})" if isinstance(x, (int, float)) else "None"
        expr = _re.sub(r"\$?([A-Z]+)\$?(\d+)", sub, body)
        if "None" in expr:
            return v
        try:
            return float(eval(expr, {"__builtins__": {}}, {}))
        except Exception:
            return v

    def g(p, _d=0):
        return _expr(a.cell(_find(am, p), 2).value, _d)

    hdr = [wb["2_Income"].cell(5, c).value for c in range(2, 60)]
    hdr = [h for h in hdr if h]
    nh = sum(1 for h in hdr if str(h).endswith("A")); ny = len(hdr) - nh
    scen = int(g("选择情景"))

    def path3(prefix, n):
        out = []
        for i in range(n):
            r = _find(am, f"{hdr[nh+i]} {prefix}")
            v = a.cell(r, 1 + scen).value
            if isinstance(v, str):          # 驱动分解接入后是跨表引用，逐层解析
                v = _cell_value(wb, a, f"{'ABCD'[scen]}{r}")
            out.append(float(v))
        return out

    o = {k: float(g(k)) for k in
         ["现金及等价物", "短期投资", "应收账款", "存货", "其他流动资产", "固定资产净额",
          "长期投资", "商誉与无形资产", "其他非流动资产", "应付账款", "其他流动负债",
          "一年内到期债务", "短期借款", "长期债务", "其他非流动负债", "少数股东权益",
          "归母权益", "夹层权益", "期初研发资本（Damodaran）",
          "基期营业收入", "摊薄股数（百万股）", "现价（美元）"]}
    gen = {k: float(g(k)) for k in
           ["研发费用 / 营收", "销售及管理费用 / 营收", "其他经营损益 / 营收", "有效所得税率",
            "少数股东损益占比", "其他流动资产 / 营收", "其他流动负债 / 营收",
            "应收周转天数 DSO", "存货周转天数 DIO", "应付周转天数 DPO",
            "折旧率（占期初固定资产净额）", "无形摊销 / 营收", "有息负债利率", "现金收益率",
            "派息率（占归母净利）", "最低现金余额", "现金归集阈值"]}
    dcf = {k: float(g(k)) for k in
           ["无风险利率 Rf（10Y UST）", "股权风险溢价 ERP", "Beta（5年对标普500回归+Blume）",
            "信用利差", "债务权重 D/(D+E)", "永续增长率 g", "终值所得税率 t_term",
            "期中折现开关（1=是 0=否）", "估值基准日滚动年数", "非经营资产（加回）",
            "少数股东扣除方式（1=账面 2=价值）", "SBC 加回开关（0=费用化 1=加回）",
            "研发资本化开关（1=开 0=关）", "研发摊销年限（年）",
            "终值 ROIC 处理（1=沿用末年 2=收敛至 WACC+溢价）",
            "终值 ROIC 收敛溢价（模式2 用）"]}
    fin = {k: float(g(k)) for k in
           ["年度回购金额 / 营收", "回购均价（美元/股）", "增发股数（百万股）", "增发价（美元/股）",
            "增发费用率", "增发到账年份序号（1=首个预测年）", "募资投向 — 资本开支", "桥接折现年数"]}
    dg = float(g("全期营收增速 Δ")); dm = float(g("全期毛利率 Δ"))
    it = None
    try:
        sw = float(g("锚定开关"))
    except KeyError:
        sw = None
    if sw is not None:
        # 首年锚定：读出各季**输入**，用与建表同一条规则函数独立重算（不读工作簿的公式结果）
        from us_calibrate import quarterly_estimate

        def quarters(who, name, n):
            out = []
            for i in range(1, n + 1):
                try:
                    out.append(float(g(f"{who}第{i}季{name}")))
                except KeyError:
                    break
            return out
        cur_r, cur_c = quarters("本年", "营业收入", 3), quarters("本年", "营业成本", 3)
        prev_r, prev_c = quarters("去年", "营业收入", 4), quarters("去年", "营业成本", 4)
        q, _why = quarterly_estimate(prev_r, prev_c, cur_r, cur_c)
        if q:
            it = dict(sw=sw, k=q["k"], rev=sum(cur_r), cogs=sum(cur_c),
                      fy_rev=q["fy_rev"], fy_cogs=q["fy_cogs"])
    return dict(hdr=hdr, nh=nh, ny=ny, o=o, gen=gen, dcf=dcf, fin=fin, dg=dg, dm=dm, it=it,
                growth=[x + dg for x in path3("营收增速", ny)],
                gm=[x + dm for x in path3("毛利率", ny)],
                capex=path3("资本开支 / 营收", ny),
                sbc=path3("SBC / 营收", ny))


def recompute(inp: dict):
    src, G, D, F = inp["o"], inp["gen"], inp["dcf"], inp["fin"]
    ny = inp["ny"]; t = G["有效所得税率"]
    o = dict(cash=src["现金及等价物"], sti=src["短期投资"], ar=src["应收账款"],
             inv=src["存货"], oca=src["其他流动资产"], ppe=src["固定资产净额"],
             lti=src["长期投资"], gwi=src["商誉与无形资产"], onca=src["其他非流动资产"],
             ap=src["应付账款"], ocl=src["其他流动负债"], cur=src["一年内到期债务"],
             std=src["短期借款"], ltd=src["长期债务"], oncl=src["其他非流动负债"],
             mi=src["少数股东权益"], eq=src["归母权益"], sh=src["摊薄股数（百万股）"],
             rda=src["期初研发资本（Damodaran）"])
    rev = src["基期营业收入"]; rows = []
    growth = list(inp["growth"]); gmp = list(inp["gm"])
    IT = inp.get("it")
    if IT and IT["sw"] == 1 and IT.get("fy_rev"):
        # 锚定：首年生效增速/毛利率由已披露季度按年初至今同比外推（与「1_Assumptions」生效值列一致）
        fy_rev, fy_cogs = IT["fy_rev"], IT["fy_cogs"]
        growth[0] = fy_rev / src["基期营业收入"] - 1
        gmp[0] = 1 - fy_cogs / fy_rev
    for i in range(ny):
        rev = rev * (1 + growth[i]); cogs = rev * (1 - gmp[i])
        gp = rev - cogs
        rd = rev * G["研发费用 / 营收"]; sga = rev * G["销售及管理费用 / 营收"]
        oo = rev * G["其他经营损益 / 营收"]
        ebit = gp - rd - sga + oo
        ie = (o["std"] + o["cur"] + o["ltd"]) * G["有息负债利率"]
        ii = (o["cash"] + o["sti"]) * G["现金收益率"]
        pbt = ebit - ie + ii
        tax = pbt * t; ni = pbt - tax
        mi = ni * G["少数股东损益占比"]; nip = ni - mi
        sbc = rev * inp["sbc"][i]
        ar = rev * G["应收周转天数 DSO"] / 365
        iv = cogs * G["存货周转天数 DIO"] / 365
        oca = rev * G["其他流动资产 / 营收"]
        ap = cogs * G["应付周转天数 DPO"] / 365
        ocl = rev * G["其他流动负债 / 营收"]
        dep = o["ppe"] * G["折旧率（占期初固定资产净额）"]
        amt = rev * G["无形摊销 / 营收"]
        fcx = (F["募资投向 — 资本开支"] / 3
               if F["募资投向 — 资本开支"] > 0
               and F["增发到账年份序号（1=首个预测年）"] <= i + 1 <= F["增发到账年份序号（1=首个预测年）"] + 2
               else 0.0)
        capex = -(rev * inp["capex"][i] + fcx)          # 负数 = 流出
        ppe = o["ppe"] - capex - dep
        gwi = o["gwi"] - amt
        dwc = -(ar - o["ar"]) - (iv - o["inv"]) - (oca - o["oca"]) + (ap - o["ap"]) + (ocl - o["ocl"])
        ocf = ni + dep + amt + sbc + dwc
        div = -nip * G["派息率（占归母净利）"]
        bb = -rev * F["年度回购金额 / 营收"]
        iss = (F["增发股数（百万股）"] * F["增发价（美元/股）"] * (1 - F["增发费用率"])
               if F["增发到账年份序号（1=首个预测年）"] == i + 1 else 0.0)
        pre = o["cash"] + ocf + capex + div + bb + iss
        mc, sw = G["最低现金余额"], G["现金归集阈值"]
        rv = (mc - pre) if pre < mc else (-min(o["std"], pre - (mc + sw))
                                          if pre > mc + sw and o["std"] > 0 else 0.0)
        cash = pre + rv; std = o["std"] + rv
        bb_sh = -bb / F["回购均价（美元/股）"]
        sbc_sh = sbc / F["回购均价（美元/股）"] if D["SBC 加回开关（0=费用化 1=加回）"] == 1 else 0.0
        iss_sh = F["增发股数（百万股）"] if F["增发到账年份序号（1=首个预测年）"] == i + 1 else 0.0
        sh = o["sh"] - bb_sh + sbc_sh + iss_sh
        eq = o["eq"] + nip + div + bb + iss + sbc
        mie = o["mi"] + mi
        ta = cash + o["sti"] + ar + iv + oca + ppe + o["lti"] + gwi + o["onca"]
        tl = std + o["cur"] + ap + o["ltd"] + ocl + o["oncl"]
        # 研发资本滚动（报表外）：摊销 = 上期资本 / 年限；资本 = 上期 + 本期研发 − 摊销
        rd_am = o["rda"] / D["研发摊销年限（年）"] if D["研发摊销年限（年）"] else 0.0
        rda = o["rda"] + rd - rd_am
        fcff = (ebit * (1 - t) + dep + amt
                + (sbc if D["SBC 加回开关（0=费用化 1=加回）"] == 1 else 0.0)
                + capex + dwc)
        rows.append(dict(rev=rev, cogs=cogs, gp=gp, ebit=ebit, pbt=pbt, tax=tax, ni=ni,
                         mi=mi, nip=nip, sbc=sbc, eps=nip / sh, sh=sh, ar=ar, inv=iv,
                         oca=oca, ap=ap, ocl=ocl, ppe=ppe, gwi=gwi, cash=cash, std=std,
                         ta=ta, tl=tl, eq=eq, mie=mie, te=eq + mie, check=ta - tl - eq - mie,
                         ocf=ocf, dep=dep, amt=amt, capex=capex, fcff=fcff,
                         nwc=ar + iv + oca - ap - ocl, ie=ie, ii=ii,
                         rda=rda, rd_am=rd_am, onca=o["onca"], oncl=o["oncl"]))
        o.update(cash=cash, ar=ar, inv=iv, oca=oca, ppe=ppe, gwi=gwi, std=std,
                 ap=ap, ocl=ocl, eq=eq, mi=mie, sh=sh, rda=rda)
    ke = D["无风险利率 Rf（10Y UST）"] + D["Beta（5年对标普500回归+Blume）"] * D["股权风险溢价 ERP"]
    kd = D["无风险利率 Rf（10Y UST）"] + D["信用利差"]
    w = (1 - D["债务权重 D/(D+E)"]) * ke + D["债务权重 D/(D+E)"] * kd * (1 - t)
    gt = D["永续增长率 g"]; tt = D["终值所得税率 t_term"]; my = D["期中折现开关（1=是 0=否）"]
    last = rows[-1]
    # 投入资本：营运资本 + 固定资产 + 商誉无形 + 其他非流动资产 − 其他非流动负债。
    # 旧式只有前两项，AAPL 会得到 −22 亿 → ROIC −4760%。
    ic_rep = last["nwc"] + last["ppe"] + last["gwi"] + last["onca"] - last["oncl"]
    ic = ic_rep + (last["rda"] if D["研发资本化开关（1=开 0=关）"] == 1 else 0.0)
    nopat = last["ebit"] * (1 - tt)
    roic_calc = nopat / ic if ic else float("inf")
    roic = ((w + D["终值 ROIC 收敛溢价（模式2 用）"])
            if D["终值 ROIC 处理（1=沿用末年 2=收敛至 WACC+溢价）"] == 2 else roic_calc)
    rr = gt / roic if roic else float("inf")
    degrade = (ic <= 0) or (roic <= w) or (rr >= 1) or (rr < 0) or not np.isfinite(rr)
    tf = nopat if degrade else last["ebit"] * (1 + gt) * (1 - tt) * (1 - rr)
    tv = nopat / w if degrade else tf / (w - gt)
    pv = sum(rows[i]["fcff"] / (1 + w) ** (i + 1 - 0.5 * my) for i in range(ny))
    pvtv = tv / (1 + w) ** (ny - 0.5 * my)
    ev0 = pv + pvtv
    rl = D["估值基准日滚动年数"]
    ev = ev0 * (1 + w) ** rl - rows[0]["fcff"] * rl * (1 + w) ** (rl / 2)
    nd = (src["短期借款"] + src["一年内到期债务"] + src["长期债务"]
          - src["现金及等价物"] - src["短期投资"])
    fin_pv = (F["增发股数（百万股）"] * F["增发价（美元/股）"] * (1 - F["增发费用率"])) \
        / (1 + w) ** F["桥接折现年数"]
    gx = ev - nd + D["非经营资产（加回）"] + fin_pv
    mi_ded = (src["少数股东权益"] if D["少数股东扣除方式（1=账面 2=价值）"] == 1
              else gx * G["少数股东损益占比"])
    eqv = gx - mi_ded
    sh0 = src["摊薄股数（百万股）"] + F["增发股数（百万股）"]
    return rows, dict(ke=ke, kd=kd, wacc=w, ic=ic, ic_rep=ic_rep, roic_calc=roic_calc,
                      roic=roic, rr=rr, nopat=nopat,
                      roic_mode=D["终值 ROIC 处理（1=沿用末年 2=收敛至 WACC+溢价）"],
                      degrade=1.0 if degrade else 0.0, tf=tf, tv=tv, pv=pv, pvtv=pvtv,
                      ev0=ev0, ev=ev, nd=nd, fin_pv=fin_pv, gx=gx, mi=mi_ded, eq=eqv,
                      ps=eqv / sh0, sh=sh0, tw=pvtv / ev0)


def verify(path: str, quiet: bool = False):
    if "%" in os.path.basename(path):
        raise ValueError(
            f"文件名含 % 字符：{os.path.basename(path)}\n"
            "formulas 引擎会把文件名塞进 %-格式模板，% 后面跟中文会抛出\n"
            "「unsupported format character」这种毫无线索的错误。请改名后重试。")
    import formulas
    inp = read_inputs(path); rows, dv = recompute(inp)
    nh, ny = inp["nh"], inp["ny"]
    CC = [CL(2 + nh + i) for i in range(ny)]
    fails = []
    wb = load_workbook(path)
    hdr = [wb["2_Income"].cell(5, c).value for c in range(2, 2 + nh + ny)]
    if len([h for h in hdr if h]) != nh + ny:
        fails.append(("列头", "不完整", hdr, None))
    if not all(str(h).endswith("E") for h in hdr[nh:]):
        fails.append(("预测列头", hdr[nh:], "应全部以 E 结尾", None))
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    bn = os.path.basename(path)

    def G(sh, cell):
        v = sol.get(f"'[{bn}]{sh.upper()}'!{cell}")
        if v is None:
            return None
        try:
            return float(v.value[0, 0])
        except Exception:
            return None

    def chk(name, sh, cell, ref):
        x = G(sh, cell)
        if x is None or (isinstance(x, float) and math.isnan(x)):
            fails.append((name, cell, "空/NaN", ref)); return
        if abs(x - float(ref)) > TOL * max(1.0, abs(float(ref))):
            fails.append((name, cell, round(x, 6), round(float(ref), 6)))

    GRID = [("营收", "2_Income", 6, "rev"), ("营业成本", "2_Income", 8, "cogs"),
            ("毛利", "2_Income", 9, "gp"), ("EBIT", "2_Income", 14, "ebit"),
            ("利息支出", "2_Income", 15, "ie"), ("利息收益", "2_Income", 16, "ii"),
            ("税前利润", "2_Income", 18, "pbt"), ("净利润", "2_Income", 20, "ni"),
            ("归母净利", "2_Income", 22, "nip"), ("摊薄股数", "2_Income", 23, "sh"),
            ("EPS", "2_Income", 24, "eps"), ("SBC", "2_Income", 26, "sbc"),
            ("现金", "3_Balance", 6, "cash"), ("应收", "3_Balance", 8, "ar"),
            ("存货", "3_Balance", 9, "inv"), ("PP&E", "3_Balance", 12, "ppe"),
            ("商誉无形", "3_Balance", 14, "gwi"), ("资产总计", "3_Balance", 16, "ta"),
            ("短期借款", "3_Balance", 18, "std"), ("负债合计", "3_Balance", 24, "tl"),
            ("归母权益", "3_Balance", 25, "eq"), ("权益合计", "3_Balance", 27, "te"),
            ("NWC", "3_Balance", 31, "nwc"), ("折旧", "4_CashFlow", 7, "dep"),
            ("摊销", "4_CashFlow", 8, "amt"), ("经营现金流", "4_CashFlow", 15, "ocf"),
            ("资本开支", "4_CashFlow", 16, "capex"), ("期末现金", "4_CashFlow", 27, "cash"),
            ("FCFF", "4_CashFlow", 19, "fcff")]
    for nm, sh, rw, key in GRID:
        for i, c in enumerate(CC):
            chk(nm, sh, f"{c}{rw}", rows[i][key])
    for i, c in enumerate(CC):
        chk("平衡校验", "3_Balance", f"{c}29", 0.0)

    dm = _rowmap(wb["5_DCF"])
    for pfx, ref in [("股权成本 Ke", dv["ke"]), ("债务成本 Kd", dv["kd"]), ("WACC", dv["wacc"]),
                     ("投入资本 IC（申报口径）", dv["ic_rep"]), ("生效投入资本 IC", dv["ic"]),
                     ("终值 ROIC（计算值）", dv["roic_calc"]), ("终值 ROIC（生效）", dv["roic"]),
                     ("稳态再投资率", dv["rr"]), ("★ 终值是否降级", dv["degrade"]),
                     ("终值 FCFF", dv["tf"]), ("终值 TV", dv["tv"]),
                     ("明确期 FCFF 现值合计", dv["pv"]), ("终值现值 PV(TV)", dv["pvtv"]),
                     ("企业价值 EV（滚动", dv["ev"]), ("减：净负债", dv["nd"]),
                     ("加：增发募资净额", dv["fin_pv"]), ("归属全体股东权益价值", dv["gx"]),
                     ("减：少数股东权益", dv["mi"]), ("归母股权价值", dv["eq"]),
                     ("当前摊薄股数", dv["sh"]), ("每股内在价值", dv["ps"]),
                     ("终值占 EV 比重", dv["tw"])]:
        try:
            chk(pfx, "5_DCF", f"B{_find(dm, pfx)}", ref)
        except KeyError:
            fails.append((pfx, "未找到行", None, ref))
    fr = _find(dm, "FCFF")
    for i in range(ny):
        chk("DCF-FCFF", "5_DCF", f"{CL(2+i)}{fr}", rows[i]["fcff"])

    # 敏感性中心格：生效 WACC + 生效 g 必须复现基准每股价值
    base_ps = G("5_DCF", f"B{_find(dm, '每股内在价值')}")
    sr = _find(dm, "WACC \\ g")
    v = G("5_DCF", f"D{sr+3}")
    if v is None or (isinstance(v, float) and math.isnan(v)):
        fails.append(("敏感性中心格", f"D{sr+3}", "空/NaN", base_ps))
    elif base_ps is None or abs(v - base_ps) > TOL * max(1.0, abs(base_ps)):
        fails.append(("敏感性中心格", f"D{sr+3}", round(v, 8), round(base_ps, 8)))

    # 生效值 ↔ 利润表 联动：假设页显示的「生效增速/毛利率」必须等于利润表实际用的那个。
    # A 股版曾出过这个 bug —— 假设页显示 30%，模型实际用 12.79%，两边都"自洽"，
    # 只有这条跨表断言能抓到。锚定把第一年从输入切换到反推，正是最容易出这种错的地方。
    GRID_RV, GRID_GM = 7, 10          # 2_Income 的「营收增速」与「毛利率」行号
    am2 = _rowmap(wb["1_Assumptions"])
    hdr_all = [wb["2_Income"].cell(5, c).value for c in range(2, 2 + nh + ny)]
    for i, c in enumerate(CC):
        yl = hdr_all[nh + i]
        for pfx, strow in ((f"{yl} 营收增速", GRID_RV), (f"{yl} 毛利率", GRID_GM)):
            try:
                er = _find(am2, pfx)
            except KeyError:
                continue
            eff = G("1_Assumptions", f"E{er}"); used = G("2_Income", f"{c}{strow}")
            if eff is None or used is None or math.isnan(eff) or math.isnan(used):
                fails.append((f"生效值联动 {pfx}", f"E{er}", "空/NaN", None))
            elif abs(eff - used) > TOL:
                fails.append((f"生效值联动 {pfx}", f"E{er} vs 2_Income!{c}{strow}",
                              round(eff, 8), round(used, 8)))
    # 敏感性引擎中心格：每个模块的中心格（Δ=0、g 与 WACC 取生效值）
    # 必须逐位复现主模型的每股价值。引擎是独立的一条计算路径 ——
    # 它与主模型走偏而无人察觉，正是"全都一致地错了"的另一种形态。
    for rr in range(1, wb["5_DCF"].max_row + 1):
        lab = wb["5_DCF"].cell(rr, 1).value
        if isinstance(lab, str) and "中心格自检" in lab:
            v = G("5_DCF", f"B{rr}")
            if v is None or (isinstance(v, float) and math.isnan(v)):
                fails.append(("引擎中心格自检", f"B{rr}", "空/NaN", 0.0))
            elif abs(v) > TOL:
                fails.append(("引擎中心格自检", f"B{rr}", round(v, 10), 0.0))
    # 杜邦恒等式：三因子/五因子连乘必须等于直接算出的 ROE。
    # 这是比率页唯一的硬断言 —— 比率之间口径打架时它会立刻暴露。
    if "7_Ratios" in wb.sheetnames:
        _wr = wb["7_Ratios"]
        for pfx in ("★ 三因子恒等式校验", "★ 五因子恒等式校验"):
            rr = next((x for x in range(1, _wr.max_row + 1)
                       if isinstance(_wr.cell(x, 1).value, str)
                       and _wr.cell(x, 1).value.startswith(pfx)), None)
            if rr is None:
                fails.append((pfx, "未找到行", None, 0.0)); continue
            found = False
            for i in range(nh + ny):
                cc = CL(2 + i); x = G("7_Ratios", f"{cc}{rr}")
                if x is None or (isinstance(x, float) and math.isnan(x)):
                    continue
                found = True
                if abs(x) > 1e-8:
                    fails.append((pfx, f"{cc}{rr}", round(x, 10), 0.0))
            if not found:
                fails.append((pfx, "全列为空", None, 0.0))
    if not quiet:
        n = len(CC) * len(GRID) + ny * 2 + 24 + (nh + ny) * 2
        print(f"验证：{'✅ 全部通过' if not fails else f'❌ {len(fails)} 项失败'}"
              f"（列头结构 + NaN 检测 + {n} 格独立复算）")
        for f in fails[:12]:
            print("   ✗", f)
    return fails, dv


def verify_all_scenarios(path: str, quiet: bool = False, values: dict | None = None):
    """三档情景全部验算 —— 只验当前情景会让 Bull/Bear 永远没被覆盖过。

    values 给定时写入各档每股价值 {"Bull": ps, "Base": ps, "Bear": ps}（与 ashare_verify
    同名同语义；ir-agent 的足球场图要用）。验算有失败的那一档也照写 —— 由调用方决定取舍。"""
    import shutil, tempfile
    out = []
    wb0 = load_workbook(path); a0 = wb0["1_Assumptions"]
    sr = _find(_rowmap(a0), "选择情景")
    cur = int(a0.cell(sr, 2).value or 2)
    for sc, nm in [(1, "Bull"), (2, "Base"), (3, "Bear")]:
        if sc == cur:
            f, dv = verify(path, quiet=True)
        else:
            # 每次一个独立临时目录：固定文件名在两次估值同时跑时会互删（测试并行时实测）
            tdir = tempfile.mkdtemp(prefix=f"_sc{sc}_")
            tmp = os.path.join(tdir, os.path.basename(path))
            try:
                shutil.copy(path, tmp)
                w = load_workbook(tmp); w["1_Assumptions"].cell(sr, 2).value = sc; w.save(tmp)
                f, dv = verify(tmp, quiet=True)
            finally:
                shutil.rmtree(tdir, ignore_errors=True)
        if values is not None and dv and dv.get("ps") is not None:
            values[nm] = float(dv["ps"])
        out += [(nm,) + x for x in f]
        if not quiet:
            print(f"   {nm}: {'✅' if not f else f'❌ {len(f)} 项'}")
    return out


def sanity(dv: dict, rows: list) -> list[str]:
    w = []
    # 高 ROIC 永续化是这类模型最容易被忽略的乐观假设：AAPL 末年 ROIC 71.8%，
    # 沿用它等于假设苹果永远赚到 71.8% 的投入资本回报，再投资率因此只有 3.5%。
    # 不替使用者做选择，但必须让他看见。
    if dv["roic"] > 0.30 and dv.get("roic_mode", 1) != 2:
        w.append(f"终值 ROIC {dv['roic']:.1%} 被永续化（再投资率仅 {dv['rr']:.1%}）："
                 f"等于假设超额回报永不被竞争侵蚀。把「终值 ROIC 处理」改为 2（收敛至 "
                 f"WACC+溢价）可看到保守口径 —— 实测四家大盘科技股因此下修 4.6%~11.3%")
    if dv["tw"] > 0.90:
        w.append(f"终值占 EV {dv['tw']:.0%}（>90%）：估值几乎全部来自终值假设")
    if dv["pv"] < 0:
        w.append(f"明确期 FCFF 现值为负（{dv['pv']:.0f}）：预测期内净消耗现金，DCF 点估值意义有限")
    if dv["degrade"] == 1.0:
        w.append(f"终值已降级为无增长永续（投入资本 {dv['ic']:.0f}，ROIC {dv['roic']:.1%} vs WACC {dv['wacc']:.2%}）"
                 "：负营运资本或低回报公司会触发，属预期行为，但意味着终值不含增长")
    if dv["eq"] <= 0:
        w.append("股权价值为负：净负债 + 少数股东权益已超过企业价值")
    neg = [i + 1 for i, r in enumerate(rows) if r["ebit"] <= 0]
    if neg:
        w.append(f"第 {neg} 个预测年 EBIT ≤ 0：亏损期不适合用 FCFF-DCF")
    return w
